"""Export a draft as LaTeX (+ .bib), Markdown, or a zip bundle with the disclosure report."""

from __future__ import annotations

import io
import json
import re
import zipfile

from . import bibtex, provenance, storage
from .models import Project

_LATEX_SPECIAL = {
    "\\": r"\textbackslash{}", "&": r"\&", "%": r"\%", "$": r"\$", "#": r"\#", "_": r"\_",
    "{": r"\{", "}": r"\}", "~": r"\textasciitilde{}", "^": r"\textasciicircum{}",
}


def latex_escape(text: str) -> str:
    return "".join(_LATEX_SPECIAL.get(ch, ch) for ch in text)


def _sentence_latex(text: str, keys: list[str]) -> str:
    body = latex_escape(text)
    if keys:
        cite = r"~\cite{" + ",".join(keys) + "}"
        m = re.search(r"([.!?])\s*$", body)
        body = body[: m.start()] + cite + m.group(1) if m else body + cite
    return body


def _paragraphs_by_section(project: Project) -> list[tuple[str, list[str]]]:
    out: list[tuple[str, list[str]]] = []
    if not project.draft:
        return out
    for p in project.draft.paragraphs:
        text = " ".join(_sentence_latex(s.text, s.ref_keys) for s in p.sentences)
        if out and out[-1][0] == p.section:
            out[-1][1].append(text)
        else:
            out.append((p.section, [text]))
    return out


def _table_latex(t) -> str:
    cols = "l" * len(t.columns)
    head = " & ".join(latex_escape(c) for c in t.columns) + r" \\"
    rows = "\n".join(" & ".join(latex_escape(v) for v in r) + r" \\" for r in t.rows)
    label = t.id.lower()
    return (f"\\begin{{table}}[t]\n\\centering\n\\caption{{{latex_escape(t.caption)}}}\n\\label{{tab:{label}}}\n"
            f"\\begin{{tabular}}{{{cols}}}\n\\hline\n{head}\n\\hline\n{rows}\n\\hline\n\\end{{tabular}}\n\\end{{table}}\n")


def to_latex(project: Project) -> str:
    L = project.ledger
    template = L.template
    if template == "ieee":
        preamble = "\\documentclass[conference]{IEEEtran}\n\\usepackage{cite}\n"
    elif template == "acm":
        preamble = "\\documentclass[sigconf]{acmart}\n"
    else:
        preamble = "\\documentclass[11pt]{article}\n\\usepackage[margin=1in]{geometry}\n\\usepackage{cite}\n"
    preamble += "\\usepackage{graphicx}\n\\usepackage[utf8]{inputenc}\n"

    authors = " \\and ".join(latex_escape(a) for a in L.authors) or "Anonymous"
    sections = _paragraphs_by_section(project)
    abstract = next((paras for name, paras in sections if name == "Abstract"), [])

    parts = [preamble, f"\\title{{{latex_escape(L.title)}}}", f"\\author{{{authors}}}"]
    if template == "acm":
        if abstract:
            parts.append("\\begin{abstract}\n" + "\n\n".join(abstract) + "\n\\end{abstract}")
        parts.append("\\begin{document}\n\\maketitle")
    else:
        parts.append("\\begin{document}\n\\maketitle")
        if abstract:
            parts.append("\\begin{abstract}\n" + "\n\n".join(abstract) + "\n\\end{abstract}")
    if L.keywords and template == "ieee":
        parts.append("\\begin{IEEEkeywords}\n" + ", ".join(latex_escape(k) for k in L.keywords) + "\n\\end{IEEEkeywords}")

    used_tables, used_figures = used_floats(project) if project.draft else (L.tables, L.figures)

    def floats() -> list[str]:
        out = [_table_latex(t) for t in used_tables]
        for f in used_figures:
            graphic = (f"\\includegraphics[width=\\linewidth]{{figures/{f.filename}}}" if f.filename
                       else "\\fbox{Figure placeholder}")
            out.append(f"\\begin{{figure}}[t]\n\\centering\n{graphic}\n\\caption{{{latex_escape(f.caption)}}}\n"
                       f"\\label{{fig:{f.id.lower()}}}\n\\end{{figure}}")
        return out

    body = [(name, paras) for name, paras in sections if name != "Abstract"]
    float_after = "Results" if any(n == "Results" for n, _ in body) else (body[-1][0] if body else None)
    for name, paras in body:
        parts.append(f"\\section{{{latex_escape(name)}}}\n" + "\n\n".join(paras))
        if name == float_after:
            parts += floats()
    if float_after is None:
        parts += floats()

    parts.append("\\section*{AI-use disclosure}\n" + latex_escape(provenance.disclosure_statement(project)))
    style = "ACM-Reference-Format" if template == "acm" else "IEEEtran" if template == "ieee" else "plain"
    parts.append(f"\\bibliographystyle{{{style}}}\n\\bibliography{{references}}\n\\end{{document}}\n")
    return "\n\n".join(parts)


def to_bib(project: Project) -> str:
    return "\n\n".join(bibtex.to_bibtex(r) for r in project.ledger.references) + "\n"


def to_markdown(project: Project) -> str:
    L = project.ledger
    lines = [f"# {L.title}", ""]
    if L.authors:
        lines += [", ".join(L.authors), ""]
    current = None
    if project.draft:
        for p in project.draft.paragraphs:
            if p.section != current:
                lines += [f"## {p.section}", ""]
                current = p.section
            sents = []
            for s in p.sentences:
                cite = f" [{'; '.join(s.ref_keys)}]" if s.ref_keys else ""
                t = s.text
                sents.append(t[:-1] + cite + t[-1] if cite and t[-1:] in ".!?" else t + cite)
            lines += [" ".join(sents), ""]
    if L.references:
        lines += ["## References", ""]
        for r in L.references:
            lines.append(f"- [{r.key}] {r.authors}. {r.title}. {r.venue} {r.year}".strip())
    return "\n".join(lines) + "\n"


def used_floats(project: Project):
    """Tables and figures linked to claims the draft actually expresses (uploaded extras stay out)."""
    L = project.ledger
    known = L.claim_map()
    expressed = {cid for p in (project.draft.paragraphs if project.draft else []) for s in p.sentences for cid in s.claim_ids}
    t_ids = {t for cid in expressed if cid in known for t in known[cid].tables}
    f_ids = {f for cid in expressed if cid in known for f in known[cid].figures}
    return [t for t in L.tables if t.id in t_ids], [f for f in L.figures if f.id in f_ids]


def to_docx(project: Project) -> bytes:
    """Word document: title block, abstract, numbered sections, numbered citations, tables, figures,
    references and the AI-use disclosure."""
    import docx
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches, Pt

    L = project.ledger
    d = docx.Document()
    normal = d.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)

    title = d.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(L.title)
    run.bold, run.font.size = True, Pt(17)
    if L.authors:
        a = d.add_paragraph(", ".join(L.authors))
        a.alignment = WD_ALIGN_PARAGRAPH.CENTER
    if L.affiliations:
        a = d.add_paragraph("; ".join(L.affiliations))
        a.alignment = WD_ALIGN_PARAGRAPH.CENTER
        a.runs[0].italic = True

    ref_no = {r.key: i for i, r in enumerate(L.references, start=1)}

    def text_of(para) -> str:
        out = []
        for s in para.sentences:
            t = s.text
            if s.ref_keys:
                cite = " [" + ", ".join(str(ref_no[k]) for k in s.ref_keys if k in ref_no) + "]"
                t = t[:-1] + cite + t[-1] if t[-1:] in ".!?" else t + cite
            out.append(t)
        return " ".join(out)

    paragraphs = project.draft.paragraphs if project.draft else []
    abstract = [p for p in paragraphs if p.section == "Abstract"]
    if abstract:
        ab = d.add_paragraph()
        ab.add_run("Abstract. ").bold = True
        ab.add_run(" ".join(text_of(p) for p in abstract))
    if L.keywords:
        kw = d.add_paragraph()
        kw.add_run("Keywords: ").bold = True
        kw.add_run(", ".join(L.keywords))

    tables, figures = used_floats(project)
    body_sections: list[str] = []
    for p in paragraphs:
        if p.section != "Abstract" and p.section not in body_sections:
            body_sections.append(p.section)
    floats_after = "Results" if "Results" in body_sections else (body_sections[-1] if body_sections else None)

    for n, name in enumerate(body_sections, start=1):
        d.add_heading(f"{n}. {name}", level=1)
        for p in paragraphs:
            if p.section == name:
                d.add_paragraph(text_of(p))
        if name == floats_after:
            for i, t in enumerate(tables, start=1):
                cap = d.add_paragraph(f"Table {i}. {t.caption}")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                grid = d.add_table(rows=1 + len(t.rows), cols=len(t.columns))
                grid.style = "Table Grid"
                for j, c in enumerate(t.columns):
                    grid.cell(0, j).text = c
                for r_i, row in enumerate(t.rows, start=1):
                    for j, v in enumerate(row[: len(t.columns)]):
                        grid.cell(r_i, j).text = v
            for i, f in enumerate(figures, start=1):
                path = storage.local_file(project.id, f.filename) if f.filename else None
                if path and path.is_file() and path.suffix.lower() in (".png", ".jpg", ".jpeg", ".gif", ".bmp"):
                    try:
                        d.add_picture(str(path), width=Inches(5.5))
                        d.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
                    except Exception:  # noqa: BLE001 - an image Word cannot embed still gets its caption
                        pass
                cap = d.add_paragraph(f"Fig. {i}. {f.caption}")
                cap.alignment = WD_ALIGN_PARAGRAPH.CENTER

    d.add_heading("AI-use disclosure", level=1)
    d.add_paragraph(provenance.disclosure_statement(project))
    if L.references:
        d.add_heading("References", level=1)
        for i, r in enumerate(L.references, start=1):
            d.add_paragraph(f"[{i}] " + ". ".join(x for x in (r.authors, r.title, r.venue, r.year) if x) + ".")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def bundle(project: Project) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("paper.tex", to_latex(project))
        z.writestr("references.bib", to_bib(project))
        z.writestr("paper.md", to_markdown(project))
        z.writestr("paper.docx", to_docx(project))
        z.writestr("ai_disclosure.md", provenance.disclosure_markdown(project))
        for f in project.ledger.figures:
            path = storage.local_file(project.id, f.filename) if f.filename else None
            if path and path.is_file():
                z.write(path, f"figures/{f.filename}")
        z.writestr("provenance.json", json.dumps({
            "stats": provenance.stats(project),
            "ledger": project.ledger.model_dump(mode="json"),
            "draft": project.draft.model_dump(mode="json") if project.draft else None,
            "history": project.history,
        }, indent=2))
    return buf.getvalue()
